import re
import os
import hashlib
from urllib.parse import urlparse, parse_qs, urlencode
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from collections import defaultdict

html_dir = '/workspace/html'

trae_variants = [
    ('text_to_image', 'https://trae-api-cn.mchost.guru/api/ide/v1/text_to_image'),
    ('text-to-image', 'https://trae-api-cn.mchost.guru/api/ide/v1/text-to-image'),
]

yukings_images = {
    'img/yukings-logo.svg': 'https://www.yukings.net/img/yukings-logo.svg',
    'img/yukings-factory.webp': 'https://www.yukings.net/img/yukings-factory.webp',
}

def reconstruct_url(img_path):
    if img_path in yukings_images:
        return yukings_images[img_path]
    
    match = re.match(r'img/(.+)-([a-f0-9]{8})\.webp$', img_path)
    if not match:
        return None
    
    prompt_part = match.group(1)
    expected_hash = match.group(2)
    
    prompt = prompt_part
    
    for variant_name, base_url in trae_variants:
        params = {'prompt': prompt, 'image_size': 'landscape_16_9'}
        url = base_url + '?' + urlencode(params)
        actual_hash = hashlib.md5(url.encode()).hexdigest()[:8]
        if actual_hash == expected_hash:
            return url
    
    return f"(需确认) {img_path}"

all_occurrences = []
img_path_pattern = re.compile(r'img/[^"\'\s)]+')

for filename in sorted(os.listdir(html_dir)):
    if not filename.endswith('.html'):
        continue
    filepath = os.path.join(html_dir, filename)
    with open(filepath, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    for line_num, line in enumerate(lines, 1):
        for match in img_path_pattern.finditer(line):
            img_path = match.group(0)
            if not img_path.endswith(('.webp', '.jpg', '.jpeg', '.png', '.gif', '.svg', '.bmp')):
                continue
            
            context = ''
            line_lower = line.lower()
            if 'background-image' in line_lower or ('background:' in line_lower and 'url(' in line_lower):
                context = 'CSS background-image (行内样式)'
            elif '<img' in line_lower and 'src=' in line_lower:
                context = 'HTML img src'
            elif 'og:image' in line_lower or 'twitter:image' in line_lower:
                context = 'Meta OG/Twitter Image'
            elif 'logo' in line_lower:
                context = 'Logo'
            elif 'json' in line_lower or 'schema' in line_lower or 'ld+' in line_lower:
                context = 'JSON-LD结构化数据'
            else:
                context = '其他'
            
            original_url = reconstruct_url(img_path)
            
            all_occurrences.append({
                'file': filename,
                'line': line_num,
                'new_path': img_path,
                'original_url': original_url,
                'context': context,
                'line_content': line.strip()[:150]
            })

print(f"Total image occurrences found: {len(all_occurrences)}")
print(f"Unique image paths: {len(set(o['new_path'] for o in all_occurrences))}")

wb = Workbook()
ws = wb.active
ws.title = "图片URL映射明细"

headers = [
    '序号',
    '所在页面',
    '代码行号',
    '图片用途类型',
    '原外部URL',
    '新相对路径',
    '代码上下文片段'
]

header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="082457", end_color="082457", fill_type="solid")
header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border = Border(
    left=Side(style='thin', color='d8dde8'),
    right=Side(style='thin', color='d8dde8'),
    top=Side(style='thin', color='d8dde8'),
    bottom=Side(style='thin', color='d8dde8')
)
orange_fill = PatternFill(start_color="FFF4EE", end_color="FFF4EE", fill_type="solid")

for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = thin_border

row_num = 2
for idx, occ in enumerate(all_occurrences, 1):
    ws.cell(row=row_num, column=1, value=idx).border = thin_border
    ws.cell(row=row_num, column=2, value=occ['file']).border = thin_border
    ws.cell(row=row_num, column=3, value=occ['line']).border = thin_border
    ws.cell(row=row_num, column=4, value=occ['context']).border = thin_border
    cell_url = ws.cell(row=row_num, column=5, value=occ['original_url'])
    cell_url.border = thin_border
    if occ['original_url'] and occ['original_url'].startswith('('):
        cell_url.fill = orange_fill
    ws.cell(row=row_num, column=6, value=occ['new_path']).border = thin_border
    ws.cell(row=row_num, column=7, value=occ['line_content']).border = thin_border
    row_num += 1

ws2 = wb.create_sheet(title="唯一图片汇总")
ws2_headers = ['序号', '新相对路径', '原外部URL', '总引用次数', '引用页面数', '所在页面列表']
for col, header in enumerate(ws2_headers, 1):
    cell = ws2.cell(row=1, column=col, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = thin_border

path_files = defaultdict(set)
path_count = defaultdict(int)
path_url = {}
for occ in all_occurrences:
    path_count[occ['new_path']] += 1
    path_files[occ['new_path']].add(occ['file'])
    path_url[occ['new_path']] = occ['original_url']

row_num2 = 2
sorted_paths = sorted(path_count.items(), key=lambda x: x[1], reverse=True)
for idx, (path, count) in enumerate(sorted_paths, 1):
    files_list = ', '.join(sorted(path_files[path]))
    ws2.cell(row=row_num2, column=1, value=idx).border = thin_border
    ws2.cell(row=row_num2, column=2, value=path).border = thin_border
    cell_url2 = ws2.cell(row=row_num2, column=3, value=path_url.get(path, ''))
    cell_url2.border = thin_border
    if path_url.get(path, '') and path_url.get(path, '').startswith('('):
        cell_url2.fill = orange_fill
    ws2.cell(row=row_num2, column=4, value=count).border = thin_border
    ws2.cell(row=row_num2, column=5, value=len(path_files[path])).border = thin_border
    ws2.cell(row=row_num2, column=6, value=files_list).border = thin_border
    row_num2 += 1

ws3 = wb.create_sheet(title="统计概览")
stats = [
    ['统计项', '数值', '备注'],
    ['扫描HTML文件总数', len([f for f in os.listdir(html_dir) if f.endswith('.html')]), '个'],
    ['图片引用总次数', len(all_occurrences), '处'],
    ['唯一图片数量', len(set(o['new_path'] for o in all_occurrences)), '张'],
    ['含图片的页面数', len(set(o['file'] for o in all_occurrences)), '个'],
    ['CSS背景图引用', len([o for o in all_occurrences if 'background' in o['context']]), '处'],
    ['HTML img标签引用', len([o for o in all_occurrences if 'img src' in o['context']]), '处'],
    ['Meta OG/社交图', len([o for o in all_occurrences if 'OG' in o['context']]), '处'],
    ['JSON-LD结构化数据', len([o for o in all_occurrences if 'JSON' in o['context']]), '处'],
    ['已成功还原原URL数', len([o for o in all_occurrences if o['original_url'] and not o['original_url'].startswith('(')]), '处'],
]

for r, row_data in enumerate(stats, 1):
    for c, val in enumerate(row_data, 1):
        cell = ws3.cell(row=r, column=c, value=val)
        cell.border = thin_border
        if r == 1:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

col_widths = {
    'A': 8,
    'B': 45,
    'C': 10,
    'D': 24,
    'E': 70,
    'F': 55,
    'G': 90,
}
for col_letter, width in col_widths.items():
    ws.column_dimensions[col_letter].width = width

ws2.column_dimensions['A'].width = 8
ws2.column_dimensions['B'].width = 55
ws2.column_dimensions['C'].width = 70
ws2.column_dimensions['D'].width = 12
ws2.column_dimensions['E'].width = 12
ws2.column_dimensions['F'].width = 100

ws3.column_dimensions['A'].width = 28
ws3.column_dimensions['B'].width = 15
ws3.column_dimensions['C'].width = 10

ws.freeze_panes = 'A2'
ws2.freeze_panes = 'A2'

output_path = os.path.join(html_dir, 'image-url-mapping.xlsx')
wb.save(output_path)
print(f"\n✅ Excel文件已生成: {output_path}")
print(f"   工作表1: 图片URL映射明细 ({len(all_occurrences)} 行)")
print(f"   工作表2: 唯一图片汇总 ({len(sorted_paths)} 行)")
print(f"   工作表3: 统计概览")

verified = len([o for o in all_occurrences if o['original_url'] and not o['original_url'].startswith('(')])
print(f"\n   已验证还原URL: {verified}/{len(all_occurrences)}")
